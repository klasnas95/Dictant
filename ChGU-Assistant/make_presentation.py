"""
ЧГУ.Ассистент — генератор презентаций (прототип).

Локальная языковая модель (Ollama) составляет план и текст слайдов,
скрипт собирает из него файл PowerPoint в оформлении ЧГУ.Ассистента.

Примеры:
    python make_presentation.py "Закон Ома" --slides 10 --audience "студенты 1 курса"
    python make_presentation.py "Закон Ома" --demo          # без модели, на готовом примере
"""
import argparse
import json
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

DARK = RGBColor(0x0E, 0x3B, 0x36)
GREEN = RGBColor(0x12, 0xA3, 0x7A)
MINT = RGBColor(0x9F, 0xD9, 0xC4)
INK = RGBColor(0x17, 0x26, 0x2B)
GREY = RGBColor(0x7B, 0x8A, 0x91)
TINT = RGBColor(0xEE, 0xF5, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

SYSTEM_PROMPT = (
    "Ты — ЧГУ.Ассистент, помощник преподавателя Чеченского государственного университета. "
    "Ты готовишь учебные презентации на русском языке: точные факты, простые формулировки, "
    "короткие пункты (до 15 слов), без выдуманных цитат и источников. "
    "Отвечай ТОЛЬКО JSON без пояснений."
)

USER_PROMPT = """Составь презентацию.
Тема: {topic}
Аудитория: {audience}
Количество слайдов с содержанием: {slides}

Формат ответа — JSON:
{{
  "title": "название презентации",
  "subtitle": "подзаголовок",
  "slides": [
    {{"title": "заголовок слайда", "bullets": ["пункт 1", "пункт 2", "пункт 3"], "notes": "что сказать преподавателю на этом слайде (2-3 предложения)"}}
  ],
  "questions": ["вопрос для самопроверки 1", "вопрос 2", "вопрос 3"]
}}
На каждом слайде 3–5 пунктов. Включи определения, формулы (если уместно), примеры и вывод."""


def ask_ollama(url, model, topic, audience, slides, timeout=600):
    """Отправляет запрос локальной модели и возвращает разобранный JSON."""
    body = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": USER_PROMPT.format(topic=topic, audience=audience, slides=slides)},
        ],
        "stream": False,
        "format": "json",
        "think": False,
        "options": {"temperature": 0.4},
    }
    try:
        reply = _post(url.rstrip("/") + "/api/chat", body, timeout)
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", "replace")
        if "think" not in detail:
            raise RuntimeError(f"Ollama вернула ошибку {e.code}: {detail}") from e
        body.pop("think")  # модель без режима рассуждений
        reply = _post(url.rstrip("/") + "/api/chat", body, timeout)
    return parse_outline(reply["message"]["content"])


def _post(url, body, timeout):
    req = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def parse_outline(text):
    """Достаёт JSON из ответа модели и приводит его к ожидаемой структуре."""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise ValueError("В ответе модели нет JSON")
    data = json.loads(text[start:end + 1])
    slides = []
    for s in data.get("slides") or []:
        if not isinstance(s, dict) or not s.get("title"):
            continue
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()][:6]
        slides.append({"title": str(s["title"]).strip(), "bullets": bullets, "notes": str(s.get("notes") or "").strip()})
    if not slides:
        raise ValueError("Модель не вернула ни одного слайда")
    return {
        "title": str(data.get("title") or "Презентация").strip(),
        "subtitle": str(data.get("subtitle") or "").strip(),
        "slides": slides,
        "questions": [str(q).strip() for q in (data.get("questions") or []) if str(q).strip()][:6],
    }


def _text(slide, text, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
    run.font.color.rgb = color
    return tf


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def _footer(slide, n):
    _text(slide, "Подготовлено с помощью ЧГУ.Ассистент", 0.6, 7.0, 6, 0.3, 10, GREY)
    _text(slide, str(n), 12.2, 7.0, 0.5, 0.3, 10, GREY, align=PP_ALIGN.RIGHT)


def _bullets(slide, items, x, y, w, h, size=26):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        marker = p.add_run()
        marker.text = "●  "
        marker.font.size, marker.font.name = Pt(size - 6), FONT
        marker.font.color.rgb = GREEN
        run = p.add_run()
        run.text = item
        run.font.size, run.font.name = Pt(size), FONT
        run.font.color.rgb = INK


def build_pptx(outline, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # Титульный слайд
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, 13.333, 7.5, DARK)
    _text(s, "ЧГУ.АССИСТЕНТ · УЧЕБНЫЕ МАТЕРИАЛЫ", 0.6, 2.5, 12, 0.5, 16, GREEN, bold=True)
    _text(s, outline["title"], 0.6, 3.1, 12, 1.4, 44, WHITE, bold=True)
    if outline["subtitle"]:
        _text(s, outline["subtitle"], 0.6, 4.5, 12, 0.9, 22, MINT)
    _text(s, "ЧГУ им. А.А. Кадырова · подготовлено с помощью ЧГУ.Ассистент", 0.6, 6.6, 12, 0.4, 14, MINT)

    # Слайды с содержанием
    n = 1
    for item in outline["slides"]:
        n += 1
        s = prs.slides.add_slide(blank)
        _rect(s, 0, 0, 13.333, 7.5, WHITE)
        _text(s, item["title"], 0.6, 0.45, 12.1, 1.0, 34, DARK, bold=True)
        _rect(s, 0.6, 1.75, 12.1, 4.95, TINT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
        _bullets(s, item["bullets"], 1.0, 2.0, 11.3, 4.5)
        _footer(s, n)
        if item["notes"]:
            s.notes_slide.notes_text_frame.text = item["notes"]

    # Вопросы для самопроверки
    if outline["questions"]:
        n += 1
        s = prs.slides.add_slide(blank)
        _rect(s, 0, 0, 13.333, 7.5, DARK)
        _text(s, "Вопросы для самопроверки", 0.6, 0.6, 12, 1.0, 36, WHITE, bold=True)
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.9), Inches(12), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, q in enumerate(outline["questions"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(14)
            r = p.add_run()
            r.text = f"{i + 1}. {q}"
            r.font.size, r.font.name = Pt(24), FONT
            r.font.color.rgb = WHITE

    prs.save(path)
    return n


def safe_name(topic):
    name = re.sub(r"[^\w\- ]+", "", topic, flags=re.U).strip().replace(" ", "_")
    return (name or "presentation")[:60] + ".pptx"


def main():
    ap = argparse.ArgumentParser(description="ЧГУ.Ассистент — генератор презентаций")
    ap.add_argument("topic", help="тема презентации")
    ap.add_argument("--slides", type=int, default=8, help="число слайдов с содержанием (3–20)")
    ap.add_argument("--audience", default="студенты 1 курса", help="для кого презентация")
    ap.add_argument("--model", default="qwen3:8b", help="модель Ollama")
    ap.add_argument("--url", default="http://localhost:11434", help="адрес сервера Ollama")
    ap.add_argument("--out", help="имя выходного файла .pptx")
    ap.add_argument("--demo", action="store_true", help="без модели: использовать demo_outline.json")
    ap.add_argument("--outline", help="готовый план презентации в JSON (используется ЧГУ.Ассистентом)")
    args = ap.parse_args()
    args.slides = max(3, min(20, args.slides))

    if args.outline:
        outline = parse_outline(Path(args.outline).read_text(encoding="utf-8-sig"))
    elif args.demo:
        outline = parse_outline((Path(__file__).parent / "demo_outline.json").read_text(encoding="utf-8"))
        print("Демо-режим: используется готовый план презентации «Закон Ома».")
    else:
        print(f"Запрос к модели {args.model}… (на слабом ПК это может занять несколько минут)")
        try:
            outline = ask_ollama(args.url, args.model, args.topic, args.audience, args.slides)
        except (urllib.error.URLError, ConnectionError) as e:
            sys.exit(f"Не удалось подключиться к Ollama по адресу {args.url}: {e}\n"
                     f"Проверьте, что Ollama запущена (команда: ollama serve) и модель скачана (ollama pull {args.model}).\n"
                     "Для проверки без модели используйте --demo.")
        except (ValueError, json.JSONDecodeError) as e:
            sys.exit(f"Модель вернула ответ в неожиданном формате: {e}. Попробуйте ещё раз или другую модель.")

    out = args.out or safe_name(outline["title"] if (args.demo or args.outline) else args.topic)
    total = build_pptx(outline, out)
    print(f"Готово: {out} — {total} слайдов (с заметками для преподавателя).")


if __name__ == "__main__":
    main()
